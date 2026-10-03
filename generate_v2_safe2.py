import docx

doc = docx.Document('Nicholas_Napoli_Resume_v2.docx')
for para in doc.paragraphs:
    for run in para.runs:
        if 'Architected Lakehouse data pipelines distributed by leveraging' in run.text:
            run.text = run.text.replace('Architected Lakehouse data pipelines distributed by leveraging', 'Architected distributed lakehouse data pipelines leveraging')

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if 'Architected Lakehouse data pipelines distributed by leveraging' in run.text:
                        run.text = run.text.replace('Architected Lakehouse data pipelines distributed by leveraging', 'Architected distributed lakehouse data pipelines leveraging')

doc.save('Nicholas_Napoli_Resume_v2.docx')
