import docx

doc = docx.Document('Nicholas_Napoli_Resume_v2.docx')
found = False
for para in doc.paragraphs:
    if 'Architected' in para.text:
        print(para.text)
        found = True

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                if 'Architected' in para.text:
                    print(para.text)
                    found = True
