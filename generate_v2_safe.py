import docx

doc = docx.Document('Nicholas_Napoli_Resume.docx')

for para in doc.paragraphs:
    for run in para.runs:
        if 'Presidio (PHI/PII Sanitization)' in run.text:
            run.text = run.text.replace('Presidio (PHI/PII Sanitization)', 'Microsoft Presidio (PHI/PII Sanitization)')
        if 'MBSE, JIRA' in run.text:
            run.text = run.text.replace('MBSE, JIRA', 'Model-Based Systems Engineering (MBSE), JIRA')

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    if 'Presidio (PHI/PII Sanitization)' in run.text:
                        run.text = run.text.replace('Presidio (PHI/PII Sanitization)', 'Microsoft Presidio (PHI/PII Sanitization)')
                    if 'MBSE, JIRA' in run.text:
                        run.text = run.text.replace('MBSE, JIRA', 'Model-Based Systems Engineering (MBSE), JIRA')

doc.save('Nicholas_Napoli_Resume_v2.docx')
