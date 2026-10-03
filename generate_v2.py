import docx
import re

doc = docx.Document('Nicholas_Napoli_Resume.docx')

for para in doc.paragraphs:
    if 'Presidio (PHI/PII Sanitization)' in para.text:
        para.text = para.text.replace('Presidio (PHI/PII Sanitization)', 'Microsoft Presidio (PHI/PII Sanitization)')
    if 'MBSE, JIRA' in para.text:
        para.text = para.text.replace('MBSE, JIRA', 'Model-Based Systems Engineering (MBSE), JIRA')

# Also check tables
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                if 'Presidio (PHI/PII Sanitization)' in para.text:
                    para.text = para.text.replace('Presidio (PHI/PII Sanitization)', 'Microsoft Presidio (PHI/PII Sanitization)')
                if 'MBSE, JIRA' in para.text:
                    para.text = para.text.replace('MBSE, JIRA', 'Model-Based Systems Engineering (MBSE), JIRA')

doc.save('Nicholas_Napoli_Resume_v2.docx')
