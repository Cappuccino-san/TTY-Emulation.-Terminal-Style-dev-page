import docx

doc = docx.Document('Nicholas_Napoli_Resume.docx')
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(f"{i}: {para.text}")

print("--- TABLES ---")
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            print(cell.text)
